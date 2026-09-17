"""
Check that the sentences in the revised manuscript still agree with the tables
they cite.

verify_numbers.py checks the analysis outputs, and it cannot see the
manuscript. That is how one-tailed P values survived in the Results after every
figure and Table S6 had already been converted: nothing compared the prose with
the table. This script closes that gap.

Each claim below pins the phrase the manuscript uses, with the number left as a
slot that is filled from the table. The phrase then has to occur verbatim in the
clean docx. If the table moves, the filled phrase stops matching; if someone
edits the sentence, the phrase stops matching. Either way this fails and names
the sentence.

Not shipped in the code release: it needs the manuscript, which the release does
not carry.

Run: python verify_manuscript_text.py
Exit status is non-zero if any check fails.
"""

from pathlib import Path
import argparse
import re
import sys

import docx
import pandas as pd

HERE = Path(__file__).resolve().parent

# Defaults are the working-tree layout. The submission archive puts the same two
# inputs elsewhere, so both are settable; run it there with
#   python verify_manuscript_text.py \
#       --docx ../01_Manuscript/Manuscript_CIR-26-0753-ET_revised.docx \
#       --tables ../04_Supplementary_Tables
ap = argparse.ArgumentParser(description=__doc__)
ap.add_argument("--docx", type=Path,
                default=HERE / "01_Main_Text" / "Manuscript_R1_clean.docx",
                help="the clean revised manuscript")
ap.add_argument("--tables", type=Path, default=HERE / "04_Tables",
                help="directory holding ST6 and ST8")
args = ap.parse_args()

CLEAN = args.docx
ST6 = args.tables / "ST6_two_sided_sensitivity.csv"
ST8 = args.tables / "ST8_nfkb_pseudobulk_sensitivity.csv"
for f in (CLEAN, ST6, ST8):
    if not f.exists():
        sys.exit(f"not found: {f}\nPass --docx and --tables for this layout.")

failures, checks = [], 0
TEXT = "\n".join(p.text for p in docx.Document(CLEAN).paragraphs)

sweep = pd.read_csv(ST6)
pseudobulk = pd.read_csv(ST8)


def two_tailed(fragment):
    rows = sweep[sweep["Comparison"].str.contains(fragment, regex=False)]
    if len(rows) != 1:
        return None
    return float(rows["P, two-tailed (revised)"].iloc[0])


def claim(label, template, *values):
    """The template, filled from the table, must appear verbatim in the text."""
    global checks
    checks += 1
    if any(v is None for v in values):
        failures.append(f"{label}: the table row it cites was not found")
        return
    phrase = template.format(*values)
    if phrase not in TEXT:
        failures.append(f"{label}: the manuscript does not contain "
                        f"“{phrase}”")


# --------------------------------------- Results against Table S6 (R1.3c)
claim("CEACAM6 pre-treatment epithelium",
      "two-sided P = {:.3f}, rank-biserial r = 0.88",
      two_tailed("CEACAM6 expression, pre-treatment"))
claim("CEACAM5 pre-treatment epithelium",
      "two-sided P = {:.3f}, r = 0.75",
      two_tailed("CEACAM5 expression, pre-treatment"))
claim("CEACAM5 in PRJEB25780",
      "CEACAM5, two-sided P = {:.3f}, r = 0.37",
      two_tailed("CEACAM5 expression, PRJEB25780"))
claim("CEACAM6 in PRJEB25780",
      "CEACAM6, two-sided P = {:.3f}, r = 0.38",
      two_tailed("CEACAM6 expression, PRJEB25780"))
claim("IHC summed score",
      "two-sided P = {:.3f}, r = 0.88, Mann–Whitney U test, n = 4 per group",
      two_tailed("CEACAM5 + CEACAM6 summed"))
claim("IHC markers separately",
      "(CEACAM5, P = {:.2f}; CEACAM6, P = {:.2f}; Table S7)",
      two_tailed("IHC staining, CEACAM5 only"),
      two_tailed("IHC staining, CEACAM6 only"))
claim("S-MP4 metaprogram",
      "two-sided exact permutation P = {:.3f}",
      two_tailed("S-MP4 score"))
claim("S-MP5 metaprogram",
      "(P = {:.3f}; Fig. 2H, I)",
      two_tailed("S-MP5 score"))
claim("neighborhood epithelial density",
      "two-sided P = {:.3f}, matched-pairs rank-biserial r = 0.85",
      two_tailed("Neighbourhood epithelial density"))
claim("distance to stroma",
      "two-sided P = {:.3f}, Wilcoxon signed-rank test; Fig. 3I",
      two_tailed("Distance to stroma"))
claim("distance to immune-rich regions",
      "P = {:.3f}; Fig. 3J",
      two_tailed("Distance to immune-rich"))
claim("BACH1 regulon",
      "two-sided exact permutation test, P = {:.3f}; Table S6)",
      two_tailed("BACH1 regulon activity"))
claim("NFKB1 regulon",
      "two-sided exact permutation test, P = {:.3f}; Fig. 5E, Table S6)",
      two_tailed("NFKB1 regulon activity"))
# Round 41 (2026-09-16): the PD-L1 sentence gives its four P values in one
# parenthesis (the "reaching significance ... approaching it" phrasing went,
# the bootstrap and BH statements are cited to Table S6), so the two claims
# above became one pinning all four table values in the order the text
# names the cell types.
claim("PD-L1 in the four cell types",
      "(P = {:.3f}, {:.3f}, {:.3f} and {:.3f}, respectively; Fig. 5J, Fig. S7, "
      "Table S6)",
      two_tailed("post-treatment monocytes/macrophages"),
      two_tailed("post-treatment epithelial cells"),
      two_tailed("post-treatment fibroblasts"),
      two_tailed("post-treatment dendritic cells"))

# ------------------------------------------ Table S8 remains the NF-κB check
checks += 1
required_pb = {"Cell type", "Timepoint", "limma-voom NES", "DESeq2 NES",
               "Per-cell NES (primary analysis; Fig. S10C)"}
missing_pb = required_pb.difference(pseudobulk.columns)
if missing_pb or len(pseudobulk) != 26:
    failures.append("Table S8 pseudobulk sensitivity is incomplete: "
                    f"missing={sorted(missing_pb)}, rows={len(pseudobulk)}")

# The immunohistochemistry is not an independent cohort, and the Results must
# not say it is; ST1 and ST5 list the same eight patients.
checks += 1
if "immunohistochemical analysis on independent samples" in TEXT:
    failures.append("the Results still describe the immunohistochemistry as "
                    "performed on independent samples; Tables S1 and S5 list the "
                    "same eight patients")

# ------------------------------------------------ no one-tailed P survives
# The Methods statement is the only place the words belong.
ALLOWED_ONE_TAILED = (
    "Statistical analyses were performed using Python",
)
checks += 1
for par in docx.Document(CLEAN).paragraphs:
    if not re.search(r"[Oo]ne-tail|[Oo]ne-sided", par.text):
        continue
    if not any(par.text.startswith(a) for a in ALLOWED_ONE_TAILED):
        failures.append("a one-tailed test is still declared outside Methods: "
                        f"“{par.text[:90]}…”")

# -------------------------------- every supplementary item cited and legended
legends = {int(m.group(1)) for m in
           (re.match(r"Figure S(\d+)\.", p.text) for p in
            docx.Document(CLEAN).paragraphs) if m}
tab_legends = {int(m.group(1)) for m in
               (re.match(r"Supplementary Table (\d+)\.", p.text) for p in
                docx.Document(CLEAN).paragraphs) if m}
cited_figs = {int(m) for m in re.findall(r"Fig(?:s?\.|ure)? ?S(\d+)", TEXT)}
cited_tabs = {int(m) for m in re.findall(r"Table S(\d+)", TEXT)}

for label, have, want in (
        # Ten supplementary figures since 2026-09-16 (S1 split into S1 and
        # S2; the former S2-S9 are S3-S10).
        ("supplementary figure legends", legends, set(range(1, 11))),
        ("supplementary table legends", tab_legends, set(range(1, 9))),
        ("supplementary figures cited", cited_figs, set(range(1, 11))),
        ("supplementary tables cited", cited_tabs, set(range(1, 9)))):
    checks += 1
    missing = sorted(want - have)
    if missing:
        failures.append(f"{label}: missing {missing}")
    extra = sorted(have - want)
    if extra:
        failures.append(f"{label}: refers to items that do not exist {extra}")

# ------------------------- every panel cited by its own label, not "S9C, D"
# A compound citation reads correctly but is invisible to any tool, editorial or
# otherwise, that searches for the panel label. Each panel gets its full label.
# Panel labels under the numbering of 2026-09-16: the former S2D/E/H are
# S3D/E/H, S3E is S4E, S7A-C are S8A-C, S8A/B are S9A/B with the coefficient
# half of the old S8B now S9C, and S9A-E are S10A-E.
PANELS = ("S3D S3E S3H S4E S8A S8B S8C S9A S9B S9C "
          "S10A S10B S10C S10D S10E").split()
first_legend = next(i for i, par in enumerate(docx.Document(CLEAN).paragraphs)
                    if re.match(r"Figure S1\.", par.text))
BODY = "\n".join(par.text for par in
                 docx.Document(CLEAN).paragraphs[:first_legend])
checks += 1
uncited = [pn for pn in PANELS if not re.search(rf"\b{pn}\b", BODY)]
if uncited:
    failures.append("panels never cited by their own label in the body text: "
                    f"{uncited}")

# ------------------------------- Figure 4B node letters against labels.py
# The panel names its nodes by NODE_LETTER (one owner, 00_Config/shared/
# labels.py); the legend must define every letter in that table's own words.
# The table lives in 00_Config beside 04_Manuscript_R1 in the working tree,
# and in the code release's 00_Config inside the submission archive, where
# this script is staged beside the manuscript. Neither found is a failure,
# not a skip.
_CONFIG = next((c for c in (HERE.parent / "00_Config",
                            HERE.parent / "06_Code" / "code" / "00_Config",
                            HERE.parent.parent / "06_Code" / "code" / "00_Config")
                if (c / "shared" / "labels.py").exists()), None)
if _CONFIG is None:
    sys.exit("shared/labels.py (the Figure 4B node-letter table) was not found "
             "beside this script or in the code release; the node-letter check "
             "cannot run")
sys.path.insert(0, str(_CONFIG))
from shared.labels import NODE_LETTER, NODE_LETTER_LEGEND  # noqa: E402
checks += 1
if NODE_LETTER_LEGEND not in TEXT:
    failures.append("Figure 4 legend does not carry the node-letter key "
                    f"“{NODE_LETTER_LEGEND}”")
for letter in sorted(set(NODE_LETTER.values())):
    checks += 1
    if not re.search(rf"\b{letter}, [a-zA-Z0-9+/ ]+?(?:;|$)", NODE_LETTER_LEGEND):
        failures.append(f"node letter {letter} is used by the panel but not "
                        f"defined in the legend key")
# MUTATION: a key with one letter dropped must not be found in the text, or
# the check above is matching something looser than the sentence.
_tokens = NODE_LETTER_LEGEND.split("; ")
_mutant = "; ".join(_tokens[:-1])
checks += 1
if _mutant + ")" in TEXT and NODE_LETTER_LEGEND not in TEXT:
    failures.append("MUTANT PASSED: a node-letter key missing its last entry "
                    "is accepted")

# ------------------------------- scatter P values against the panels' files
# Since 2026-09-14 (evening) the scatter panels (Figure 3 A, B, D, E; Figure 5
# K, M) print rho and leave P to the legend. The panel scripts write their
# statistics through 00_Config/cnsfig/corr_stats.py; the legend sentence is
# generated from the same files by edits.py, and is checked here against
# them again, so a stale docx or a re-run panel cannot drift apart quietly.
import importlib.util as _ilu
_spec = _ilu.spec_from_file_location("corr_stats", _CONFIG / "cnsfig" / "corr_stats.py")
_corr = _ilu.module_from_spec(_spec); _spec.loader.exec_module(_corr)
for fig, letters in (("3", "ABDE"), ("5", "KM")):
    want = _corr.sentence(fig, letters)
    body = want.split(": ", 1)[1]          # the "(A) x, P; ... ." part
    checks += 1
    if body not in TEXT:
        failures.append(f"Figure {fig} legend does not carry the scatter P "
                        f"values the panels wrote: \u201c{body}\u201d")
    # MUTATION: the same sentence with one P value changed must not be in
    # the text; if it is, the check is matching a looser pattern.
    rows = _corr.read(fig, letters[0])
    mutant = body.replace(rows[0]["p_printed"], "P = 0.42", 1)
    checks += 1
    if mutant == body or mutant in TEXT:
        failures.append(f"MUTANT PASSED: Figure {fig} legend accepts a wrong "
                        f"scatter P value")

print(f"Checked {checks} claims in the manuscript against the tables it cites.")
if failures:
    print(f"\n{len(failures)} FAILED:\n")
    for f in failures:
        print(f"  - {f}")
    sys.exit(1)
print("Every sentence agrees with its table.")
